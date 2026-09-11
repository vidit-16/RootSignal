"""Write analytical results into readable Excel workbooks.

This is a presentation layer and nothing more. Every figure it writes was
computed and tested upstream; no metric is derived here. If a number in a
workbook disagreed with the same number in the mart, the workbook would be the
thing that is wrong.

Spreadsheets get forwarded, split up and pasted into decks, and they arrive
stripped of whatever context surrounded them. Each workbook therefore opens with
a sheet stating what it contains, when it was produced, and what its figures do
and do not mean — so an estimate cannot travel onward as a measurement.
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="595959")

CURRENCY_FORMAT = "#,##0.00"
PERCENT_FORMAT = "0.0%"
RATE_FORMAT = "0.000"
INTEGER_FORMAT = "#,##0"

MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10

# Columns whose name ends with one of these takes that number format. Matching
# on meaning rather than on dtype keeps a rate from being shown as currency
# merely because it happens to be stored as a float.
SUFFIX_FORMATS = {
    "_pct": PERCENT_FORMAT,
    "_share": PERCENT_FORMAT,
    "_rate": RATE_FORMAT,
    "attainment": PERCENT_FORMAT,
    "_units": INTEGER_FORMAT,
    "_count": INTEGER_FORMAT,
    "_sales": CURRENCY_FORMAT,
    "_target": CURRENCY_FORMAT,
    "_value": CURRENCY_FORMAT,
}
EXACT_FORMATS = {
    "net_sales": CURRENCY_FORMAT,
    "gross_sales": CURRENCY_FORMAT,
    "impact": CURRENCY_FORMAT,
    "variance": CURRENCY_FORMAT,
    "contribution": CURRENCY_FORMAT,
    "aov": CURRENCY_FORMAT,
    "units": INTEGER_FORMAT,
    "orders": INTEGER_FORMAT,
    "wape": RATE_FORMAT,
    "mae": CURRENCY_FORMAT,
    "rmse": CURRENCY_FORMAT,
    "movement": RATE_FORMAT,
    "priority_score": CURRENCY_FORMAT,
}


def _number_format(column: str) -> str | None:
    name = str(column).lower()
    if name in EXACT_FORMATS:
        return EXACT_FORMATS[name]
    for suffix, fmt in SUFFIX_FORMATS.items():
        if name.endswith(suffix):
            return fmt
    return None


def _fit_columns(sheet: Worksheet, frame: pd.DataFrame) -> None:
    for position, column in enumerate(frame.columns, start=1):
        longest = max(
            [len(str(column))] + [len(str(value)) for value in frame[column].head(200)]
        )
        width = min(MAX_COLUMN_WIDTH, max(MIN_COLUMN_WIDTH, longest + 2))
        sheet.column_dimensions[get_column_letter(position)].width = width


def write_table(
    workbook: Workbook,
    name: str,
    frame: pd.DataFrame,
    note: str | None = None,
) -> Worksheet:
    """Write one frame to its own sheet, formatted for reading rather than parsing.

    An empty frame still produces a sheet carrying its column headers, so a
    reader can tell the difference between a report that found nothing and a
    report that failed to run.
    """
    # Excel sheet names are capped at 31 characters and reject several symbols.
    safe_name = str(name)[:31]
    for character in "[]:*?/\\":
        safe_name = safe_name.replace(character, "-")

    sheet = workbook.create_sheet(safe_name)
    start_row = 1

    if note:
        sheet.cell(row=1, column=1, value=note).font = NOTE_FONT
        start_row = 3

    for position, column in enumerate(frame.columns, start=1):
        cell = sheet.cell(row=start_row, column=position, value=str(column))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_offset, (_, row) in enumerate(frame.iterrows(), start=start_row + 1):
        for position, column in enumerate(frame.columns, start=1):
            value = row[column]
            if isinstance(value, (list, tuple, dict)):
                value = "; ".join(str(item) for item in value) if value else ""
            elif pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.date()
            cell = sheet.cell(row=row_offset, column=position, value=value)
            number_format = _number_format(column)
            if number_format and isinstance(value, (int, float)):
                cell.number_format = number_format

    if len(frame):
        sheet.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(frame.columns))}{start_row + len(frame)}"
        )
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)
    _fit_columns(sheet, frame)
    return sheet


def write_cover(
    workbook: Workbook,
    title: str,
    description: str,
    notes: Sequence[str] = (),
    generated_at: datetime | None = None,
) -> Worksheet:
    """Open the workbook with what it is and what its figures mean.

    Placed first deliberately. A sheet of numbers with no statement of its basis
    invites an estimate being quoted as a measurement once the file has been
    forwarded a few times.
    """
    sheet = workbook.create_sheet("Read me", 0)
    sheet.cell(row=1, column=1, value=title).font = TITLE_FONT
    stamp = (generated_at or datetime.now()).strftime("%Y-%m-%d %H:%M")
    sheet.cell(row=2, column=1, value=f"Generated {stamp}").font = NOTE_FONT

    sheet.cell(row=4, column=1, value=description).alignment = Alignment(wrap_text=True)

    row = 6
    if notes:
        sheet.cell(row=row, column=1, value="Please read before using these figures").font = Font(
            bold=True
        )
        row += 1
        for note in notes:
            cell = sheet.cell(row=row, column=1, value=f"•  {note}")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

    sheet.column_dimensions["A"].width = 110
    return sheet


def new_workbook() -> Workbook:
    """An empty workbook with openpyxl's default sheet removed."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def save(workbook: Workbook, destination: str | Path) -> Path:
    """Write the workbook, creating the directory if needed."""
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not workbook.sheetnames:
        raise ValueError("Refusing to save a workbook with no sheets.")
    workbook.save(path)
    return path
