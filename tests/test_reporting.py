from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

from rootsignal.analysis import summarise_by_period
from rootsignal.reporting import (
    available_reports,
    build_all_reports,
    build_report,
    new_workbook,
    save,
    write_cover,
    write_table,
)


# --------------------------------------------------------------------------
# Workbook mechanics
# --------------------------------------------------------------------------


def test_an_empty_result_still_produces_a_sheet_with_headers(tmp_path) -> None:
    """A reader must be able to tell "found nothing" from "failed to run"."""
    workbook = new_workbook()
    write_table(workbook, "Signals", pd.DataFrame(columns=["segment", "impact"]))
    path = save(workbook, tmp_path / "empty.xlsx")

    sheet = load_workbook(path)["Signals"]
    assert [cell.value for cell in sheet[1]] == ["segment", "impact"]


def test_sheet_names_are_made_legal_for_excel(tmp_path) -> None:
    """Excel caps names at 31 characters and rejects several symbols."""
    workbook = new_workbook()
    write_table(workbook, "Accuracy/net_sales [weekly] " + "x" * 20, pd.DataFrame({"a": [1]}))
    path = save(workbook, tmp_path / "names.xlsx")

    name = load_workbook(path).sheetnames[0]
    assert len(name) <= 31
    assert not set(name) & set("[]:*?/\\")


def test_number_formats_follow_what_a_column_means(tmp_path) -> None:
    """A rate stored as a float is not currency merely because it is a float."""
    workbook = new_workbook()
    write_table(
        workbook,
        "Metrics",
        pd.DataFrame({"net_sales": [1234.5], "fill_rate": [0.93], "variance_pct": [-0.1]}),
    )
    path = save(workbook, tmp_path / "formats.xlsx")
    sheet = load_workbook(path)["Metrics"]

    assert sheet.cell(row=2, column=1).number_format == "#,##0.00"
    assert sheet.cell(row=2, column=2).number_format == "0.000"
    assert sheet.cell(row=2, column=3).number_format == "0.0%"


def test_cover_sheet_carries_its_notes(tmp_path) -> None:
    workbook = new_workbook()
    write_table(workbook, "Data", pd.DataFrame({"a": [1]}))
    write_cover(workbook, "A title", "A description", notes=["First caveat", "Second caveat"])
    path = save(workbook, tmp_path / "cover.xlsx")

    workbook = load_workbook(path)
    assert workbook.sheetnames[0] == "Read me"
    text = " ".join(
        str(cell.value) for row in workbook["Read me"].iter_rows() for cell in row if cell.value
    )
    assert "A title" in text
    assert "First caveat" in text and "Second caveat" in text


def test_an_empty_workbook_is_not_saved(tmp_path) -> None:
    with pytest.raises(ValueError, match="no sheets"):
        save(new_workbook(), tmp_path / "nothing.xlsx")


# --------------------------------------------------------------------------
# Reports
# --------------------------------------------------------------------------


def test_every_report_builds_and_opens_with_its_caveats(cleaned_dataset, tmp_path) -> None:
    """Spreadsheets get forwarded stripped of context, so each carries its own."""
    written = build_all_reports(cleaned_dataset.tables, output_dir=tmp_path)

    assert len(written) == len(available_reports()) == 6
    for path in written:
        assert path.exists() and path.stat().st_size > 0
        workbook = load_workbook(path)
        assert workbook.sheetnames[0] == "Read me", path.name
        assert len(workbook.sheetnames) >= 2, path.name


def test_unknown_report_names_are_rejected(cleaned_dataset, tmp_path) -> None:
    with pytest.raises(ValueError, match="Unknown report"):
        build_report("quarterly_board_pack", cleaned_dataset.tables, tmp_path)


def test_tracker_figures_match_the_analytical_layer(cleaned_dataset, tmp_path) -> None:
    """Excel reports, it does not compute. A figure here must match its source.

    If a number in a workbook disagreed with the same number in the mart, the
    workbook would be the thing that is wrong.
    """
    tables = cleaned_dataset.tables
    path = build_report("daily_sales_tracker", tables, tmp_path)

    rows = list(load_workbook(path)["Daily tracker"].iter_rows(values_only=True))
    header, first = rows[0], dict(zip(rows[0], rows[1]))
    source = summarise_by_period(tables, period="day").iloc[0]

    assert "date" in header
    assert first["net_sales"] == pytest.approx(source["net_sales"])
    assert first["order_count"] == source["order_count"]
    assert first["fill_rate"] == pytest.approx(source["fill_rate"])


def test_signal_report_publishes_the_impact_assumptions(cleaned_dataset, tmp_path) -> None:
    """An impact estimate must not travel onward as a measured loss."""
    path = build_report(
        "root_signal_report",
        cleaned_dataset.tables,
        tmp_path,
        current_period="2026-02-23",
        comparison_period="2026-02-09",
    )
    workbook = load_workbook(path)
    cover = " ".join(
        str(cell.value) for row in workbook["Read me"].iter_rows() for cell in row if cell.value
    )

    assert "not a measured loss" in cover
    assert "CONSISTENT WITH" in cover
    assert "not a probability" in cover
    assert {"Signals", "Evidence", "Confidence criteria", "Impact"} <= set(workbook.sheetnames)


def test_signal_report_carries_the_reasoning_behind_each_signal(cleaned_dataset, tmp_path) -> None:
    """The evidence and criteria travel with the conclusion, not separately."""
    path = build_report(
        "root_signal_report",
        cleaned_dataset.tables,
        tmp_path,
        current_period="2026-02-23",
        comparison_period="2026-02-09",
    )
    workbook = load_workbook(path)

    signals = list(workbook["Signals"].iter_rows(values_only=True))
    criteria = list(workbook["Confidence criteria"].iter_rows(values_only=True))
    evidence = list(workbook["Evidence"].iter_rows(values_only=True))

    assert len(signals) > 1
    # Six named criteria are recorded for every signal reported.
    assert len(criteria) - 1 == (len(signals) - 1) * 6
    assert len(evidence) > 1


def test_forecast_report_covers_every_metric(cleaned_dataset, tmp_path) -> None:
    path = build_report("forecast_report", cleaned_dataset.tables, tmp_path)
    names = set(load_workbook(path).sheetnames)

    for metric in ("net_sales", "units", "orders"):
        assert f"Accuracy {metric}" in names
        assert f"Backtest {metric}" in names


def test_fill_rate_report_leaves_undefined_rates_blank(cleaned_dataset, tmp_path) -> None:
    """A period with no demand has no fill rate, and zero would be a claim."""
    path = build_report("fill_rate_report", cleaned_dataset.tables, tmp_path)
    sheet = load_workbook(path)["Service detail"]

    header = [cell.value for cell in sheet[1]]
    fill_rate_column = header.index("fill_rate")
    ordered_column = header.index("ordered_units")

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[ordered_column] == 0:
            assert row[fill_rate_column] is None
